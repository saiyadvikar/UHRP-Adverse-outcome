"""
generate_styled_excel.py  --  Adverse Outcome Excel Report Generator
====================================================================
Generates all styled Excel reports from bmo_dashboard_data.json.
Auto-called at the end of Generate_Dashboard.py.

OUTPUT
------
Excel Reports/
    Overall_Adverse_Outcomes_All_Months.xlsx
    Monthly/
        <Month>_<Year>_Adverse_Outcomes.xlsx  (one per admission month)

BLOCK SCORECARD COLUMNS (A-K)
  A Block | B Total Cases | C Eclampsia | D Preeclampsia | E Anemia (Severe)
  F Sickle cell crisis | G Stillbirth / IUFD | H Neonatal Death
  I Maternal Death | J SNCU transfer | K HRP Flagged

PATIENT LINELIST COLUMNS (A-P)
  A Sr No | B Assigned Block | C Assigned Subcentre | D Village / Area
  E Patient Name | F Mobile Number | G MP ID (RCH ID)
  H Date of Admission | I Admission Month | J DH Diagnosis
  K Clinical Category | L HRP flagged in ANMOL | M HRP High Risk Category
  N Perinatal Outcome | O Mother Status | P SNCU Transfer
"""

import json, os
from openpyxl import Workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict

# ── PATHS ────────────────────────────────────────────────────────────────
BASE = os.path.dirname(os.path.abspath(__file__))
JSON_PATH = os.path.join(BASE, "bmo_dashboard_data.json")

print("Loading data from bmo_dashboard_data.json...")
with open(JSON_PATH, encoding="utf-8") as f:
    DATA = json.load(f)
cases = DATA["cases"]
for i, c in enumerate(cases, 1):
    c.setdefault("sr_no", i)
print(f"  Loaded {len(cases)} cases.")

# ── STYLE HELPERS ─────────────────────────────────────────────────────────
def mkfill(h): return PatternFill("solid", fgColor=h)
def mkfont(h, bold=False, size=10, name="Calibri"):
    return Font(color=h, bold=bold, size=size, name=name)
def mkborder(color="BFBFBF"):
    s = Side(border_style="thin", color=color)
    return Border(left=s, right=s, top=s, bottom=s)

header_fill    = mkfill("1F4E79"); header_font    = mkfont("FFFFFF", bold=True)
cell_border    = mkborder()
fill_zebra     = mkfill("EBF3FB"); fill_white     = mkfill("FFFFFF")
fill_live      = mkfill("E2EFDA"); font_live      = mkfont("276221", bold=True)
fill_stillbirth= mkfill("FCE4D6"); font_stillbirth= mkfont("C00000", bold=True)
fill_neodeath  = mkfill("F4CCCC"); font_neodeath  = mkfont("880000", bold=True)
fill_matdeath  = mkfill("F4CCCC"); font_matdeath  = mkfont("7B0000", bold=True)
fill_referral  = mkfill("FFF2CC"); font_referral  = mkfont("7F6000", bold=True)
fill_hrp_yes   = mkfill("D6E4F0"); font_hrp_yes   = mkfont("154360", bold=True)
fill_hrp_no    = mkfill("F5F5F5"); font_hrp_no    = mkfont("666666")
fill_sncu_yes  = mkfill("FCE4D6"); font_sncu_yes  = mkfont("843C0C", bold=True)
fill_sncu_no   = mkfill("EAF4EB"); font_sncu_no   = mkfont("1E6631")
fill_preeclamp = mkfill("D9E1F2"); font_preeclamp = mkfont("1F3864", bold=True)
fill_eclamp    = mkfill("F2E0D9"); font_eclamp    = mkfont("833C00", bold=True)
fill_anemia    = mkfill("EDEDED"); font_anemia    = mkfont("595959", bold=True)
fill_scd       = mkfill("F9E4D1"); font_scd       = mkfont("7E4A0B", bold=True)
fill_out_block = mkfill("D9D9D9"); font_out_block = mkfont("595959", bold=True)
font_regular   = mkfont("000000"); font_bold      = mkfont("000000", bold=True)
font_mono      = Font(name="Courier New", size=9)

SC_FILLS = {
    "Eclampsia":         (mkfill("F2E0D9"), mkfont("833C00", bold=True)),
    "Preeclampsia":      (mkfill("D9E1F2"), mkfont("1F3864", bold=True)),
    "Anemia (Severe)":   (mkfill("EDEDED"), mkfont("595959", bold=True)),
    "Sickle cell crisis":(mkfill("F9E4D1"), mkfont("7E4A0B", bold=True)),
    "Maternal Death":    (mkfill("F4CCCC"), mkfont("7B0000", bold=True)),
    "SNCU transfer":     (mkfill("FCE4D6"), mkfont("843C0C", bold=True)),
    "Stillbirth / IUFD":(mkfill("FCE4D6"),  mkfont("C00000", bold=True)),
    "Neonatal Death":    (mkfill("F4CCCC"), mkfont("880000", bold=True)),
    "HRP Flagged":       (mkfill("D6E4F0"), mkfont("154360", bold=True)),
}

# Block Scorecard headers — 11 columns A-K
SC_HEADERS = [
    "Block", "Total Cases", "Eclampsia", "Preeclampsia", "Anemia (Severe)",
    "Sickle cell crisis", "Stillbirth / IUFD", "Neonatal Death",
    "Maternal Death", "SNCU transfer", "HRP Flagged"
]

# Patient Linelist headers — 16 columns A-P
LL_HEADERS = [
    "Sr No", "Assigned Block", "Assigned Subcentre", "Village / Area",
    "Patient Name", "Mobile Number", "MP ID (RCH ID)",
    "Date of Admission", "Admission Month", "DH Diagnosis", "Clinical Category",
    "HRP flagged in ANMOL", "HRP High Risk Category",
    "Perinatal Outcome", "Mother Status", "SNCU Transfer"
]


def build_scorecard_sheet(ws, filtered_cases):
    """Write the Block Scorecard sheet."""
    block_data = defaultdict(lambda: {
        "total": 0, "maternal_death": 0, "stillbirth": 0, "neo_death": 0,
        "eclampsia": 0, "preeclampsia": 0, "anemia": 0, "sickle": 0,
        "sncu": 0, "hrp": 0
    })
    for c in filtered_cases:
        if not c.get("is_khargone_district", True):
            continue
        b   = c["assigned_block"]; d = block_data[b]
        ms  = c.get("mother_status_norm", "")
        bs  = c.get("baby_status_norm", "")
        cat = c.get("clinical_category", "")
        d["total"] += 1
        if ms == "Maternal Death":                     d["maternal_death"] += 1
        if bs == "Stillbirth / IUFD":                  d["stillbirth"] += 1
        if "Neonatal Death" in bs:                     d["neo_death"] += 1
        if "Eclampsia" in cat and "Pre" not in cat:   d["eclampsia"] += 1
        if "Preeclampsia" in cat:                      d["preeclampsia"] += 1
        if "Anemia" in cat:                            d["anemia"] += 1
        if "Sickle" in cat or "SCD" in cat or "sickle" in str(c.get("diagnosis_raw", "")).lower() or "scd" in str(c.get("diagnosis_raw", "")).lower(): d["sickle"] += 1
        if c.get("is_sncu_admission"):                 d["sncu"] += 1
        if c.get("hrp_matched"):                       d["hrp"] += 1

    ws.append(SC_HEADERS); ws.row_dimensions[1].height = 30
    for cn in range(1, len(SC_HEADERS) + 1):
        cell = ws.cell(row=1, column=cn)
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    for ri, blk_name in enumerate(sorted(block_data.keys()), start=2):
        d = block_data[blk_name]; rf = fill_zebra if ri % 2 == 0 else fill_white
        row_vals = [
            blk_name, d["total"],
            d["eclampsia"], d["preeclampsia"], d["anemia"], d["sickle"],
            d["stillbirth"], d["neo_death"],
            d["maternal_death"], d["sncu"], d["hrp"]
        ]
        ws.append(row_vals); ws.row_dimensions[ri].height = 22
        for ci, val in enumerate(row_vals, start=1):
            cell = ws.cell(row=ri, column=ci); cell.border = cell_border
            cell.alignment = Alignment(vertical="center",
                                       horizontal="center" if ci > 1 else "left")
            hdr = SC_HEADERS[ci - 1]
            if ci == 1:
                cell.font = font_bold; cell.fill = rf
            elif val and val > 0 and hdr in SC_FILLS:
                cell.fill = SC_FILLS[hdr][0]; cell.font = SC_FILLS[hdr][1]
            else:
                cell.fill = rf; cell.font = font_regular if val == 0 else font_bold
    for col in ws.columns:
        ws.column_dimensions[get_column_letter(col[0].column)].width = 17
    ws.freeze_panes = "A2"


def populate_linelist_sheet(ws, sheet_cases):
    """Write a patient linelist sheet — 16 columns A-P."""
    ws.views.sheetView[0].showGridLines = True
    ws.append(LL_HEADERS); ws.row_dimensions[1].height = 28
    for cn in range(1, len(LL_HEADERS) + 1):
        cell = ws.cell(row=1, column=cn)
        cell.fill = header_fill; cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = cell_border

    for ri, c in enumerate(sheet_cases, start=2):
        rf        = fill_zebra if ri % 2 == 0 else fill_white
        blk       = c.get("assigned_block", "")
        cat       = c.get("clinical_category", "")
        baby_st   = c.get("baby_status_norm", "")
        mother_st = c.get("mother_status_norm", "")
        sncu_st   = ("Yes (" + c["sncu_reason"] + ")") if c.get("is_sncu_admission") else "No"

        row_vals = [
            c.get("sr_no", ri - 1), blk,
            c.get("assigned_subcentre", ""), c.get("assigned_village", ""),
            c.get("patient_name", ""),  c.get("mobile_number", ""),
            c.get("mpid", ""),
            c.get("admission_date", ""), c.get("admission_month", ""),
            c.get("diagnosis_raw", ""), cat,
            "Flagged High Risk" if c.get("hrp_matched") else "Not Flagged in HRP",
            c.get("hrp_flagged_category", ""),
            baby_st, mother_st, sncu_st
        ]
        ws.append(row_vals); ws.row_dimensions[ri].height = 24

        for ci in range(1, len(row_vals) + 1):
            cell = ws.cell(row=ri, column=ci)
            cell.fill = rf; cell.font = font_regular
            cell.border = cell_border; cell.alignment = Alignment(vertical="center")

            if ci == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_bold
            elif ci == 2:
                cell.font = font_bold
                if blk == "Out of District":
                    cell.fill = fill_out_block; cell.font = font_out_block
            elif ci == 5:
                cell.font = font_bold
            elif ci in [6, 7]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = font_mono
            elif ci in [8, 9]:
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif ci in [10, 11]:
                if "Preeclampsia" in cat:  cell.fill = fill_preeclamp; cell.font = font_preeclamp
                elif "Eclampsia" in cat:   cell.fill = fill_eclamp;    cell.font = font_eclamp
                elif "Anemia" in cat:      cell.fill = fill_anemia;    cell.font = font_anemia
                elif "Sickle" in cat:     cell.fill = fill_scd;       cell.font = font_scd
            elif ci == 12:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c.get("hrp_matched"): cell.fill = fill_hrp_yes; cell.font = font_hrp_yes
                else:                    cell.fill = fill_hrp_no;  cell.font = font_hrp_no
            elif ci == 14:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if baby_st == "Live Birth":
                    cell.fill = fill_live;       cell.font = font_live
                elif baby_st == "Stillbirth / IUFD":
                    cell.fill = fill_stillbirth; cell.font = font_stillbirth
                elif "Neonatal Death" in baby_st:
                    cell.fill = fill_neodeath;   cell.font = font_neodeath
            elif ci == 15:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if mother_st == "Maternal Death":
                    cell.fill = fill_matdeath; cell.font = font_matdeath
                elif "Referral" in mother_st or "Near-Miss" in mother_st:
                    cell.fill = fill_referral; cell.font = font_referral
                elif "Survived" in mother_st:
                    cell.fill = fill_live;     cell.font = font_live
            elif ci == 16:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                if c.get("is_sncu_admission"): cell.fill = fill_sncu_yes; cell.font = font_sncu_yes
                else:                          cell.fill = fill_sncu_no;  cell.font = font_sncu_no

    for col in ws.columns:
        cl = get_column_letter(col[0].column)
        mx = max((len(str(cell.value or "")) for cell in col), default=8)
        ws.column_dimensions[cl].width = min(max(mx + 3, 12), 42)
    ws.freeze_panes = "A2"
    if sheet_cases:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(LL_HEADERS))}{len(sheet_cases) + 1}"


def build_workbook(filtered_cases):
    """Build a complete workbook: Scorecard + block tabs + Out of District."""
    wb = Workbook(); wb.remove(wb.active)
    build_scorecard_sheet(wb.create_sheet(title="Block Scorecard"), filtered_cases)
    for blk_name in sorted(set(
        c["assigned_block"] for c in filtered_cases if c.get("is_khargone_district", True)
    )):
        blk_cases = [c for c in filtered_cases if c["assigned_block"] == blk_name]
        populate_linelist_sheet(wb.create_sheet(title=blk_name[:31]), blk_cases)
    out_cases = [c for c in filtered_cases if not c.get("is_khargone_district", True)]
    populate_linelist_sheet(wb.create_sheet(title="Out of District Inflow"), out_cases)
    return wb


# ── GENERATE REPORTS ──────────────────────────────────────────────────────
report_root = os.path.join(BASE, "Excel Reports")
monthly_dir = os.path.join(report_root, "Monthly")
os.makedirs(report_root, exist_ok=True)
os.makedirs(monthly_dir,  exist_ok=True)

print("\nBuilding Overall Compile File (all cases)...")
overall_wb   = build_workbook(cases)
overall_path = os.path.join(report_root, "Overall_Adverse_Outcomes_All_Months.xlsx")
overall_wb.save(overall_path)
print("  Saved: Overall_Adverse_Outcomes_All_Months.xlsx")

months = sorted(set(c.get("admission_month", "") for c in cases if c.get("admission_month")))
print(f"\nBuilding Monthly Files for: {months}")
for month in months:
    month_cases = [dict(c) for c in cases if c.get("admission_month") == month]
    for i, mc in enumerate(month_cases, 1):
        mc["sr_no"] = i
    month_wb   = build_workbook(month_cases)
    safe_name  = month.replace(" ", "_")
    month_path = os.path.join(monthly_dir, f"{safe_name}_Adverse_Outcomes.xlsx")
    month_wb.save(month_path)
    n = len(month_cases)
    print(f"  Saved: {safe_name}_Adverse_Outcomes.xlsx  ({n} cases)")

print("\nExcel report generation completed successfully!")
